r"""The active-syndicate denominator (market_active.py; FIX4 scope items 2, 3 and 4).

The official lists for 2020-2024 in data/market_active_syndicates.json lacked Syndicate 33 (Hiscox) in every year,
while the extraction's workbook built from Lloyd's official files lists it: the workbook minus the file was {33}
each year. The coverage figure and the systemic-share note typed their own copies of the counts, so repairing the
file alone would not have reached them. These tests hold the file to the workbook, every consumer to the one module,
the appendix's account of the off-list records to what their filings state, and the provenance note's count of
filings without a dual-model output to the loader's.

Run:  python -m pytest src/test_market_active.py -q
"""
import io
import json
import os

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import openpyxl  # noqa: E402
import pytest  # noqa: E402

import build_current_results as bcr  # noqa: E402
import generate_data_audit as gda  # noqa: E402
import make_paper_figures as MPF  # noqa: E402
import market_active  # noqa: E402

HERE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


@pytest.fixture(scope="module")
def workbook():
    return openpyxl.load_workbook(str(market_active.WORKBOOK), read_only=True, data_only=True)


def _sheet(wb, year):
    return {int(r[0]) for r in list(wb[str(year)].iter_rows(values_only=True))[1:] if isinstance(r[0], (int, float))}


def _overview(wb):
    """{year: (SFCR active count, the file's own list count)} from the workbook's Overview sheet."""
    out = {}
    for r in wb["Overview"].iter_rows(values_only=True):
        cells = [x for x in r if x is not None]
        if len(cells) >= 3 and isinstance(cells[0], int) and 2014 <= cells[0] <= 2024:
            out[cells[0]] = (cells[1], cells[2])
    return out


def test_the_official_lists_are_the_workbooks_sheets(workbook):
    lists = market_active.official_lists()
    assert sorted(lists) == [2020, 2021, 2022, 2023, 2024]
    for year, listed in lists.items():
        assert listed == _sheet(workbook, year), year


def test_syndicate_33_is_listed_every_year():
    assert all(33 in listed for listed in market_active.official_lists().values())


def test_the_counts_are_the_workbooks_own(workbook):
    """The annual reports' counts before 2020 are the workbook's SFCR column; the official lists' lengths are its
    count of each year's list (98/92/93/95/95), which the file missing Syndicate 33 fell one short of."""
    overview = _overview(workbook)
    active = market_active.active_by_year()
    for year in range(2014, 2020):
        assert active[year] == market_active.MARKET_AR[year] == overview[year][0], year
    for year in range(2020, 2025):
        assert active[year] == overview[year][1], year
    assert market_active.active_total() == sum(active.values())


def test_every_coverage_figure_reads_the_one_denominator(monkeypatch, tmp_path):
    """A list with 50 syndicates a year must reach the coverage figure's bars, the provenance note's coverage and
    the data audit: none of them may carry its own copy."""
    fake = tmp_path / "lists.json"
    fake.write_text(json.dumps({str(y): list(range(1, 51)) for y in range(2020, 2025)}), encoding="utf-8")
    monkeypatch.setattr(market_active, "LIST", fake)
    yr = np.array([2020, 2021, 2024])
    fig, ax, _legend = MPF.coverage_figure(yr, {2020: 1, 2021: 1, 2024: 1})
    heights = [p.get_height() for p in ax.containers[0]]
    plt.close(fig)
    assert heights[-5:] == [50] * 5 and heights[:6] == [market_active.MARKET_AR[y] for y in range(2014, 2020)]
    assert bcr._active_by_year()[2022] == 50
    assert gda.compute()["market"][2022] == 50


def test_the_systemic_share_note_counts_the_active_years():
    import check_systemic_share as css
    assert css.N_ACTIVE == market_active.active_total()


# ---- scope item 3: the off-list records' account ---------------------------------------------------------------

REGISTER = {"1884_2022": {"category": "WHOLE"}, "1110_2024": {"category": "PART"}, "780_2020": {"category": "AFTER"},
            "1884_2023": {"category": "NOTCOUNT", "note": "Ambiguous, so not counted. A legacy vehicle.",
                          "evidence": "Syndicate 1884 underwrites legacy reinsurance, supported by capital."},
            "1884_2024": {"category": "NOTCOUNT", "note": "A legacy vehicle that did no deal in the year.",
                          "evidence": "Syndicate 1884 underwrites legacy reinsurance, supported by capital."},
            "3268_2021": {"category": "NOTCOUNT", "note": "Contradicted by the same filing: template text.",
                          "evidence": "Whilst the Syndicate has been placed into run-off, it will continue."}}
REVIEWED = {"1110_2023": {"reason": "The filing does not say the syndicate is in run-off."}}
REGISTER_PATH = "`pdf_extraction/audit/runoff_corpus_register.json`"


def test_the_off_list_account_is_what_the_filings_state_class_by_class():
    """FIX4 scope item 3, on E's final register: a statement of run-off (WHOLE, PART, AFTER), a statement the
    register does not count (NOTCOUNT: here legacy vehicles), a filing it reviewed and found to state none, and a
    record it holds no entry for, each in its own words; nothing it does not count is called run-off."""
    text = gda.offlist_sentence(["1110_2023", "1110_2024", "1884_2022", "1884_2023", "1884_2024", "2014_2020"],
                                REGISTER, REVIEWED)
    assert text == (
        "Of the 6 corpus records whose syndicate is not on that year's list, 2 have filings that state that the "
        "syndicate is in run-off or has ceased underwriting (1110/2024, 1884/2022); 2 are legacy vehicles whose "
        "filings, on the register's reading, do not state that the syndicate is in run-off (1884/2023, 1884/2024; "
        "each says \"Syndicate 1884 underwrites legacy reinsurance ...\"); 1 has a filing that the register reviewed "
        "and found not to state that the syndicate is in run-off (1110/2023); and for 1 the register holds no "
        "entry (2014/2020). The register is " + REGISTER_PATH + ".")
    assert "run-off years" not in text and "run-off syndicates" not in text


def test_a_notcount_entry_without_the_legacy_mark_is_not_called_a_legacy_vehicle():
    text = gda.offlist_sentence(["1884_2023", "3268_2021"], REGISTER, REVIEWED)
    assert text == (
        "Of the 2 corpus records whose syndicate is not on that year's list, none has a filing that states that the "
        "syndicate is in run-off or has ceased underwriting; and 2 have filings that, on the register's reading, do "
        "not state that the syndicate is in run-off (1884/2023, 3268/2021; 1884/2023: \"Syndicate 1884 underwrites "
        "legacy reinsurance ...\"; 3268/2021: \"Whilst the Syndicate has been placed into run-off ...\"). The "
        "register is " + REGISTER_PATH + ".")
    assert "legacy vehicle" not in text


def test_the_off_list_account_counts_one_record_and_none():
    assert gda.offlist_sentence(["780_2020"], REGISTER, REVIEWED) == (
        "Of the 1 corpus record whose syndicate is not on that year's list, 1 has a filing that states that the "
        "syndicate is in run-off or has ceased underwriting (780/2020). The register is " + REGISTER_PATH + ".")
    assert gda.offlist_sentence(["1884_2024"], REGISTER, REVIEWED).startswith(
        "Of the 1 corpus record whose syndicate is not on that year's list, none has a filing that states that the "
        "syndicate is in run-off or has ceased underwriting; and 1 is a legacy vehicle whose filing, on the "
        "register's reading, does not state that the syndicate is in run-off (1884/2024; it says \"Syndicate 1884")
    assert gda.offlist_sentence(["2014_2020"], REGISTER).startswith(
        "Of the 1 corpus record whose syndicate is not on that year's list, none has a filing that states that the "
        "syndicate is in run-off or has ceased underwriting; and for 1 the register holds no entry (2014/2020)")
    assert gda.offlist_sentence([], REGISTER) == "Every corpus record's syndicate is on that year's list."


def test_the_final_registers_off_list_account():
    """The six off-list records at E's register as imported at 2ee4007e (as at 51bf5095 for these six; the import
    before the current one, 57b4b14d): 1110/2020 and 1884/2022 (WHOLE, kept in the RITC regime), 1110/2024 (PART),
    1884/2023 and 1884/2024 (NOTCOUNT: legacy vehicles whose filings do not state run-off) and 1110/2023 (reviewed,
    not run-off), read from the committed register."""
    text = gda.offlist_sentence(["1110_2020", "1110_2023", "1110_2024", "1884_2022", "1884_2023", "1884_2024"],
                                gda.load_runoff_corpus_register(), gda.load_reviewed_not_runoff())
    assert text == (
        "Of the 6 corpus records whose syndicate is not on that year's list, 3 have filings that state that the "
        "syndicate is in run-off or has ceased underwriting (1110/2020, 1110/2024, 1884/2022); 2 are legacy vehicles "
        "whose filings, on the register's reading, do not state that the syndicate is in run-off (1884/2023, "
        "1884/2024; each says \"Syndicate 1884 (“the Syndicate”) underwrites Reinsurance to Close "
        "(“RITC”) and legacy reinsurance in the Lloyd’s market ...\"); and 1 has a filing that the "
        "register reviewed and found not to state that the syndicate is in run-off (1110/2023). The register is "
        + REGISTER_PATH + ".")


def test_the_appendix_prints_the_off_list_account(monkeypatch):
    """The appendix said the off-list records "are run-off syndicates that still file accounts" (FIX4 scope item 3):
    it now prints what their filings state, from the register it reads."""
    c, r = gda.compute(), gda.mine_raw()
    monkeypatch.setattr(gda, "load_runoff_corpus_register", lambda: REGISTER)
    monkeypatch.setattr(gda, "load_reviewed_not_runoff", lambda: REVIEWED)
    text = gda.md(c, r)
    assert "  " + gda.offlist_sentence(c["offlist"], REGISTER, REVIEWED) in text
    assert "run-off syndicates that still file accounts" not in text


def test_the_off_list_records_are_the_corpus_less_the_lists():
    c = gda.compute()
    lists = market_active.official_lists()
    obs = json.load(io.open(os.path.join(HERE, "model", "exposure_results.json"), encoding="utf-8"))["observations"]
    expected = sorted("%d_%d" % (o["syndicate"], o["year"]) for o in obs
                      if o["year"] in lists and o["syndicate"] not in lists[o["year"]])
    assert sorted(c["offlist"]) == expected
    assert sum(e["extra"] for e in c["diff"].values()) == len(expected)


# ---- scope item 4: the provenance note's count of filings without a dual-model output ---------------------------

def test_the_provenance_note_counts_the_loaders_filings_without_a_dual_model_output():
    ex = json.load(io.open(os.path.join(HERE, "model", "exposure_results.json"), encoding="utf-8"))
    doc = io.open(os.path.join(HERE, "docs", "data-provenance.md"), encoding="utf-8").read()
    fresh = bcr.provenance_clauses(doc, ex)
    n = ex["disposition_flow"]["files_without_dual_model_record_overlapping_audit_count"]
    assert "The %d filings with no usable dual-model output cannot all be labelled OCR failures" % n in fresh
    planted = fresh.replace("The %d filings with no usable" % n, "The 128 filings with no usable")
    assert bcr.provenance_clauses(planted, ex) == fresh
